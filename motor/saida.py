"""Geração do Excel de recomendações (somente camada 0-20)."""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .leitor import Laudo
from .parametros import Ajustes, Parametros
from .processar import COL_SUB, colunas_doses, colunas_recomendacao, resumo_por_talhao
from .regras import fatores_calcario

FONTE = "Arial"
VERDE = "104A2A"
LARANJA = "E25A10"
CLARO = "E8F1EA"
AMARELO = "FFF4CC"
fino = Side(style="thin", color="C9D3CC")
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)
ASSETS = Path(__file__).resolve().parent.parent / "assets"


def cabecalho_marca(ws, titulo: str, subtitulo: str, n_cols: int) -> int:
    """Logo + título nas primeiras linhas. Devolve a linha onde a tabela começa."""
    logo = ASSETS / "logo.png"
    if logo.exists():
        img = XLImage(str(logo))
        img.height = 48
        img.width = int(48 * 1433 / 442)
        ws.add_image(img, "A1")
    col_txt = 4
    ws.cell(1, col_txt, titulo).font = Font(name=FONTE, bold=True, size=14, color=VERDE)
    ws.cell(2, col_txt, subtitulo).font = Font(name=FONTE, size=10, color="555555")
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 20
    for c in range(1, max(n_cols, 8) + 1):
        ws.cell(4, c).fill = PatternFill("solid", fgColor=LARANJA)
    ws.row_dimensions[4].height = 4
    # impressão: paisagem, ajustado à largura da página
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    return 6


def _tabela(ws, df: pd.DataFrame, linha0: int = 1, larguras: dict | None = None, destaque=(),
            congelar: bool = True):
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(linha0, j, col)
        c.font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
    for i, row in enumerate(df.itertuples(index=False), linha0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(i, j, v)
            col = df.columns[j - 1]
            c.font = Font(name=FONTE, size=10, bold=col in destaque)
            c.border = BORDA
            if col in destaque:
                c.fill = PatternFill("solid", fgColor=CLARO)
            if isinstance(v, float):
                c.number_format = "#,##0" if col in destaque or "kg/ha" in col else "0.0#"
    ws.row_dimensions[linha0].height = 32
    if congelar:
        ws.freeze_panes = ws.cell(linha0 + 1, 1)
    if len(df):
        ws.auto_filter.ref = f"A{linha0}:{get_column_letter(df.shape[1])}{linha0 + len(df)}"
    for j, col in enumerate(df.columns, 1):
        w = (larguras or {}).get(col)
        if w is None:
            amostra = [len(str(col)) * 0.9] + [len(str(x)) for x in df.iloc[:200, j - 1]]
            w = min(max(amostra) + 2, 60)
        ws.column_dimensions[get_column_letter(j)].width = max(w, 10)


def descrever_ajustes(res: pd.DataFrame, aj: Ajustes) -> list[tuple[str, str]]:
    linhas = []
    for nome, r in res.attrs.get("ajustes", {}).items():
        a = r.get("ajuste")
        if nome == "KCl parcelado":
            linhas.append(("KCl", f"parcelado em 2 aplicações: {r['pct_1']:g}% + {r['pct_2']:g}% da dose total"))
            continue
        if a is None:
            f = r["formula"]
            linhas.append(("Produto fosfatado", f"formulação {int(f[0]):02d}-{int(f[1]):02d}-{int(f[2]):02d}"
                                                f" · média {r['media_final']:,.0f} kg/ha".replace(",", ".")))
            continue
        if not a.ativo():
            continue
        if a.media_alvo is not None:
            txt = f"ajustado à média-alvo de {a.media_alvo:,.0f} kg/ha (cliente já comprou)"
        else:
            txt = f"ajuste de {a.pct:+g}%"
        txt += f" · média pela regra {r['media_regra']:,.0f} → {r['media_final']:,.0f} kg/ha " \
               f"(fator ×{r['fator']:.3f})"
        linhas.append((nome, txt.replace(",", ".")))
    if aj.gesso_por_s_elementar:
        linhas.append(("Gesso", "substituído por S elementar (S = 4,1904 · Gesso^0,3754)"))
    return linhas


def _laudo_original(wo, laudo: Laudo, linhas_ok: set, arquivo: str = "") -> None:
    orig = laudo.original                       # linha i ↔ linha Excel (laudo.linha_cab + 2 + i)
    r_out = 1
    if arquivo:
        wo.cell(1, 1, f"Arquivo: {arquivo}").font = Font(name=FONTE, bold=True, size=10, color=VERDE)
        r_out = 2
    for j, col in enumerate(orig.columns, 1):
        wo.cell(r_out, j, col).font = Font(name=FONTE, bold=True, size=10)
    r_out += 1
    for i, row in enumerate(orig.itertuples(index=False)):
        linha_excel = laudo.linha_cab + 2 + i
        antes_dos_dados = laudo.linha_cab + 1 + i < laudo.inicio      # ex.: linha de unidades
        if antes_dos_dados or linha_excel in linhas_ok:
            for j, v in enumerate(row, 1):
                wo.cell(r_out, j, None if (isinstance(v, float) and pd.isna(v)) else v).font = \
                    Font(name=FONTE, size=10)
            r_out += 1


def gerar_excel(res: pd.DataFrame, laudo: Laudo, par: Parametros, nome_laudo: str = "",
                ajustes: Ajustes | None = None) -> bytes:
    aj = ajustes or Ajustes()
    rec_all = res[~res[COL_SUB]].reset_index(drop=True)
    n_sub = int(res[COL_SUB].sum())
    wb = Workbook()

    fazenda = next((x for x in rec_all["Propriedade"] if x), "")
    produtor = next((x for x in rec_all["Proprietário"] if x), "")
    sub = " · ".join(x for x in (produtor, fazenda, datetime.now().strftime("%d/%m/%Y")) if x)

    # 1) Recomendações
    ws = wb.active
    ws.title = "Recomendações"
    cols = colunas_recomendacao(rec_all)
    rec = rec_all[cols]
    l0 = cabecalho_marca(ws, "Recomendação de corretivos e fertilizantes (camada 0-20 cm)", sub, len(cols))
    _tabela(ws, rec, linha0=l0, destaque=colunas_doses(rec),
            larguras={"Observações": 60, "Proprietário": 28})
    j_obs = cols.index("Observações") + 1
    for i in range(l0 + 1, l0 + 1 + len(rec)):
        c = ws.cell(i, j_obs)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if c.value:
            c.fill = PatternFill("solid", fgColor=AMARELO)
    notas = descrever_ajustes(res, aj)
    lin = l0 + len(rec) + 2
    if notas:
        ws.cell(lin, 1, "Ajustes aplicados").font = Font(name=FONTE, bold=True, size=10, color=VERDE)
        for k, (a, b) in enumerate(notas, 1):
            ws.cell(lin + k, 1, f"{a}: {b}").font = Font(name=FONTE, size=10)

    # 2) Resumo por talhão
    wr = wb.create_sheet("Resumo por talhão")
    tab = resumo_por_talhao(res)
    l0r = cabecalho_marca(wr, "Resumo por talhão", sub, tab.shape[1])
    _tabela(wr, tab, linha0=l0r)

    # 3) Memória de cálculo
    mem_cols = [c for c in rec_all.columns
                if c not in cols + [COL_SUB] or c in ("Arquivo", "N Lab", "Identificação", "Talhão", "Abertura")]
    _tabela(wb.create_sheet("Memória de cálculo"), rec_all[mem_cols])

    # 4) Parâmetros usados
    wp = wb.create_sheet("Parâmetros")
    f = fatores_calcario(par.calcario)
    info = [
        ("Laudo", nome_laudo or "-"),
        ("Aba lida", laudo.aba),
        ("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Amostras na recomendação (0-20 cm)", len(rec_all)),
        ("Amostras subsuperficiais excluídas", n_sub),
        ("Calcário usado", f"CaO {par.calcario.cao:g}% · MgO {par.calcario.mgo:g}% · PRNT {par.calcario.prnt:g}"),
        ("Fatores de correção do calcário",
         f"Ca ×{f['ca']:.3f} · Mg ×{f['mg']:.3f} · PRNT ×{f['prnt']:.3f}"),
    ]
    info += [(f"Ajuste – {a}", b) for a, b in notas]
    info += [("Aviso de leitura", a) for a in laudo.avisos]
    info += [("Coluna lida: " + k, v) for k, v in laudo.colunas.items()]
    for i, (a, b) in enumerate(info, 1):
        wp.cell(i, 1, a).font = Font(name=FONTE, bold=True, size=10)
        wp.cell(i, 2, b).font = Font(name=FONTE, size=10)
    tabp = pd.DataFrame(par.como_tabela(), columns=["Grupo", "Parâmetro", "Valor"])
    tabp["Valor"] = tabp["Valor"].astype(str)
    _tabela(wp, tabp, linha0=len(info) + 2, congelar=False)
    wp.column_dimensions["A"].width = 38
    wp.column_dimensions["B"].width = 80

    # 5) Laudo original (somente as linhas da camada 0-20) — uma aba por arquivo quando juntados
    origens = laudo.origens or [(nome_laudo, laudo)]
    for n_or, (nome_or, lo) in enumerate(origens, 1):
        titulo = "Laudo original" if len(origens) == 1 else f"Laudo original ({n_or})"
        wo = wb.create_sheet(titulo)
        if len(origens) > 1:
            sel = rec_all["Arquivo"] == nome_or if "Arquivo" in rec_all else slice(None)
            linhas_ok = set(rec_all.loc[sel, "Linha no laudo"])
        else:
            linhas_ok = set(rec_all["Linha no laudo"])
        _laudo_original(wo, lo, linhas_ok, nome_or if len(origens) > 1 else "")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
