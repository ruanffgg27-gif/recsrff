"""Excel de interpretação (extra): situação de cada talhão por classes de teores."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import interpretacao as it
from .leitor import Laudo
from .processar import colunas_doses_totais
from .saida import BORDA, FONTE, VERDE, cabecalho_marca


def _hex(c: str) -> str:
    return c.lstrip("#").upper()


def _cab(ws, linha, cols):
    for j, t in enumerate(cols, 1):
        c = ws.cell(linha, j, t)
        c.font = Font(name=FONTE, bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
    ws.row_dimensions[linha].height = 42


def _tabela_diagnostico(ws, medias: pd.DataFrame, linha0: int, atencao: dict | None) -> int:
    chaves = [k for k in it.ATRIBUTOS if k in medias]
    cols = ["Talhão", "Amostras", "Textura", "Argila (%)"] + \
           [f"{it.ATRIBUTOS[k][0]}{' ↓' if k in it.INVERTIDOS else ''}\n{it.ATRIBUTOS[k][1]}".strip()
            for k in chaves]
    if atencao is not None:
        cols.append("Pontos de atenção")
    _cab(ws, linha0, cols)
    for i, (_, r) in enumerate(medias.iterrows(), linha0 + 1):
        vals = [r["Talhão"], int(r["Amostras"]), r["Textura"],
                None if pd.isna(r["argila"]) else round(r["argila"] / 10, 1)]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name=FONTE, size=9, bold=j == 1)
            c.alignment = Alignment(vertical="center", horizontal="left" if j in (1, 3) else "center")
            c.border = BORDA
        for j, k in enumerate(chaves, 5):
            v = r[k]
            cl = it.classificar(k, v)
            c = ws.cell(i, j, None if pd.isna(v) else round(float(v), 2))
            c.number_format = "0.00" if k in ("b", "k", "al") else "0.0"
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDA
            c.font = Font(name=FONTE, size=9, bold=True,
                          color="FFFFFF" if cl in it.TEXTO_CLARO else "1A1A1A")
            if cl:
                c.fill = PatternFill("solid", fgColor=_hex(it.CORES[cl]))
        if atencao is not None:
            itens = atencao.get(r["Talhão"], [])
            c = ws.cell(i, len(cols), "; ".join(itens) if itens else "Sem pontos críticos")
            c.font = Font(name=FONTE, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            c.border = BORDA
        ws.row_dimensions[i].height = 30
    widths = [22, 9, 13, 9] + [9.5] * len(chaves) + ([48] if atencao is not None else [])
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    return linha0 + len(medias) + 1


def _legenda(ws, linha: int) -> int:
    ws.cell(linha, 1, "Legenda das classes").font = Font(name=FONTE, bold=True, size=9, color=VERDE)
    for j, cl in enumerate(it.CLASSES, 2):
        c = ws.cell(linha, j, cl)
        c.fill = PatternFill("solid", fgColor=_hex(it.CORES[cl]))
        c.font = Font(name=FONTE, size=8, bold=True, color="FFFFFF" if cl in it.TEXTO_CLARO else "1A1A1A")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(linha + 1, 1, "Valores = média das amostras do talhão. ↓ = menor valor é melhor (H+Al, Al, m%).").font = \
        Font(name=FONTE, size=8, italic=True, color="666666")
    return linha + 2


def gerar_excel_interpretacao(res: pd.DataFrame, laudo: Laudo, nome_laudo: str = "") -> bytes:
    wb = Workbook()
    d = laudo.dados
    fazenda = next((x for x in d["propriedade"] if x), "")
    produtor = next((x for x in d["proprietario"] if x), "")
    sub = " · ".join(x for x in (produtor, fazenda, datetime.now().strftime("%d/%m/%Y")) if x)

    medias = it.medias_por_talhao(laudo, subsuperficial=False)
    atencao = it.pontos_de_atencao(medias) if not medias.empty else {}

    # 1) Diagnóstico 0-20
    ws = wb.active
    ws.title = "Diagnóstico 0-20 cm"
    l0 = cabecalho_marca(ws, "Interpretação da análise de solo – camada 0-20 cm", sub, 12)
    if medias.empty:
        ws.cell(l0, 1, "Sem amostras da camada 0-20 cm neste laudo.")
    else:
        fim = _tabela_diagnostico(ws, medias, l0, atencao)
        _legenda(ws, fim + 1)
        ws.freeze_panes = ws.cell(l0 + 1, 2)

    # 2) Camada subsuperficial, se houver
    medias_sub = it.medias_por_talhao(laudo, subsuperficial=True)
    if not medias_sub.empty:
        wsub = wb.create_sheet("Camada 20-40 cm")
        l1 = cabecalho_marca(wsub, "Interpretação – camada subsuperficial (20-40 cm)", sub, 12)
        fim = _tabela_diagnostico(wsub, medias_sub, l1, it.pontos_de_atencao(medias_sub))
        _legenda(wsub, fim + 1)
        wsub.cell(fim + 3, 1, "Camada informativa (Ca, Al, m% e S em profundidade orientam o uso de gesso); "
                              "não entra no cálculo das doses.").font = Font(name=FONTE, size=8, italic=True)

    # 3) Gráficos
    wg = wb.create_sheet("Gráficos")
    lg = cabecalho_marca(wg, "Gráficos", sub, 12)
    figs = []
    if not medias.empty:
        figs.append(it.figura_png(it.figura_mapa(medias, "Situação por talhão – camada 0-20 cm")))
    doses = colunas_doses_totais(res)
    if doses and (~res["Camada subsuperficial"]).any():
        figs.append(it.figura_png(it.figura_doses(res, doses)))
    if not medias.empty:
        figs.append(it.figura_png(it.figura_panorama(medias)))
        figs.append(it.figura_png(it.figura_classes_amostras(laudo)))
        if all(c in medias for c in ("ca", "mg", "k", "hal")):
            figs.append(it.figura_png(it.figura_bases(medias)))
    if not medias_sub.empty:
        figs.append(it.figura_png(it.figura_mapa(medias_sub, "Camada subsuperficial (20-40 cm)")))
    lin = lg
    for png in figs:
        img = XLImage(io.BytesIO(png))
        escala = 1100 / img.width
        img.width, img.height = int(img.width * escala), int(img.height * escala)
        wg.add_image(img, f"A{lin}")
        lin += int(img.height / 20) + 2

    # 4) Classes em formato de lista (para filtrar)
    wl = wb.create_sheet("Classes (lista)")
    _cab(wl, 1, ["Camada", "Talhão", "Atributo", "Unidade", "Média", "Classe"])
    r = 2
    for camada, m in (("0-20 cm", medias), ("20-40 cm", medias_sub)):
        for _, row in m.iterrows():
            for k in it.ATRIBUTOS:
                if k not in m or pd.isna(row[k]):
                    continue
                cl = it.classificar(k, row[k])
                vals = [camada, row["Talhão"], it.ATRIBUTOS[k][0], it.ATRIBUTOS[k][1], round(float(row[k]), 2), cl]
                for j, v in enumerate(vals, 1):
                    c = wl.cell(r, j, v)
                    c.font = Font(name=FONTE, size=9)
                    c.border = BORDA
                c = wl.cell(r, 6)
                c.fill = PatternFill("solid", fgColor=_hex(it.CORES[cl]))
                c.font = Font(name=FONTE, size=9, bold=True, color="FFFFFF" if cl in it.TEXTO_CLARO else "1A1A1A")
                r += 1
    for j, w in enumerate([10, 22, 14, 12, 10, 14], 1):
        wl.column_dimensions[get_column_letter(j)].width = w
    wl.freeze_panes = "A2"
    wl.auto_filter.ref = f"A1:F{max(r - 1, 1)}"

    # 5) Faixas de referência
    wf = wb.create_sheet("Faixas de referência")
    _cab(wf, 1, ["Atributo", "Unidade"] + it.CLASSES)
    for j, cl in enumerate(it.CLASSES, 3):
        wf.cell(1, j).fill = PatternFill("solid", fgColor=_hex(it.CORES[cl]))
        wf.cell(1, j).font = Font(name=FONTE, size=9, bold=True,
                                  color="FFFFFF" if cl in it.TEXTO_CLARO else "1A1A1A")
    for i, (k, (rot, uni, faixas)) in enumerate(it.ATRIBUTOS.items(), 2):
        wf.cell(i, 1, rot + (" (menor é melhor)" if k in it.INVERTIDOS else "")).font = Font(name=FONTE, size=9, bold=True)
        wf.cell(i, 2, uni).font = Font(name=FONTE, size=9)
        for j, (_, lo, hi) in enumerate(faixas, 3):
            c = wf.cell(i, j, f"{lo:g} – {hi:g}".replace(".", ","))
            c.font = Font(name=FONTE, size=9)
            c.alignment = Alignment(horizontal="center")
            c.border = BORDA
    wf.column_dimensions["A"].width = 24
    wf.column_dimensions["B"].width = 12
    for j in range(3, 11):
        wf.column_dimensions[get_column_letter(j)].width = 12
    n = len(it.ATRIBUTOS) + 3
    wf.cell(n, 1, "Fonte: tabela de legendas da equipe técnica. Classes texturais pela argila (Embrapa): "
                  "arenosa < 15%; média 15–35%; argilosa 35–60%; muito argilosa ≥ 60% de argila.").font = \
        Font(name=FONTE, size=8, italic=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
